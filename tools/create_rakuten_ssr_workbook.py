from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[1]

OUTPUT = (
    ROOT
    / "data"
    / "intraday"
    / "rakuten_rss_ssr.xlsx"
)

MAX_STOCKS = 50

HEADERS = [
    "Code",
    "Name",
    "PrevClose",
    "Open",
    "OpenTime",
    "CurrentPrice",
    "CurrentTime",
    "Low",
    "LowTime",
]


def main():
    OUTPUT.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SSR"

    for col, header in enumerate(
        HEADERS,
        start=1,
    ):
        ws.cell(
            row=1,
            column=col,
            value=header,
        )

    for row in range(
        2,
        MAX_STOCKS + 2,
    ):
        ws.cell(
            row=row,
            column=2,
            value=(
                f'=RssMarket(A{row}&".T",'
                f'"銘柄名称")'
            ),
        )

        rss_items = [
            (3, "前日終値"),
            (4, "始値"),
            (5, "始値詳細時刻"),
            (6, "現在値"),
            (7, "現在値詳細時刻"),
            (8, "安値"),
            (9, "安値詳細時刻"),
        ]

        for col, item in rss_items:
            ws.cell(
                row=row,
                column=col,
                value=(
                    f'=RssMarket(A{row}&".T",'
                    f'"{item}")'
                ),
            )

    widths = {
        "A": 12,
        "B": 32,
        "C": 14,
        "D": 14,
        "E": 18,
        "F": 16,
        "G": 18,
        "H": 14,
        "I": 18,
    }

    for col, width in widths.items():
        ws.column_dimensions[col].width = width

    ws.freeze_panes = "A2"

    wb.save(OUTPUT)

    print("=" * 60)
    print(" Rakuten RSS SSR workbook created")
    print("=" * 60)
    print(f"FILE       : {OUTPUT}")
    print(f"MAX STOCKS : {MAX_STOCKS}")
    print(f"COLUMNS    : {len(HEADERS)}")


if __name__ == "__main__":
    main()
