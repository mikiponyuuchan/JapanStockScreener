from pathlib import Path
import argparse
import re
import sys
import time

import pandas as pd
import requests
from pypdf import PdfReader
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"

sys.path.insert(
    0,
    str(SRC),
)

from screener.loader import load_stock_list


def normalize_code(value):
    value = str(value).strip()

    if (
        len(value) == 5
        and value.endswith("0")
    ):
        return value[:-1]

    return value


def normalize_text(text):
    # Financial PDFs often write negative values as
    # triangle + whitespace + number, for example:
    #   U+25B3 42.6
    # Normalize that form before generic replacements
    # so the sign is not separated from the number.
    text = re.sub(
        r"[\u25b3\u25b2]\s*(?=\d)",
        "-",
        text,
    )

    replacements = {
        "\u3000": " ",
        "\u25b3": "-",
        "\u2212": "-",
        "\uff0d": "-",
        "\uff05": "%",
    }

    for old, new in replacements.items():
        text = text.replace(
            old,
            new,
        )

    text = re.sub(
        r"[ \t]+",
        " ",
        text,
    )

    return text


def to_float(value):
    if value is None:
        return None

    value = str(value).strip()
    value = value.replace(",", "")
    value = value.replace("%", "")

    try:
        return float(value)
    except Exception:
        return None


def download_pdf(
    url,
    date_key,
    code,
):
    out_dir = (
        Path("data")
        / "analysis"
        / "earnings_pdf"
        / date_key
    )

    out_dir.mkdir(
        parents=True,
        exist_ok=True,
    )

    path = out_dir / (
        f"{code}.pdf"
    )

    if path.exists():
        return path

    r = requests.get(
        url,
        timeout=30,
        headers={
            "User-Agent":
                "JapanStockScreener/1.0"
        },
    )

    r.raise_for_status()

    if (
        "pdf"
        not in str(
            r.headers.get(
                "content-type",
                "",
            )
        ).lower()
    ):
        raise RuntimeError(
            "response is not PDF"
        )

    path.write_bytes(
        r.content
    )

    time.sleep(0.15)

    return path


def extract_first_page(path):
    reader = PdfReader(
        str(path)
    )

    if not reader.pages:
        return ""

    text = (
        reader.pages[0]
        .extract_text()
        or ""
    )

    return normalize_text(
        text
    )


def find_performance_row(text):
    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    candidates = []

    for line in lines:
        if (
            "\u9023\u7d50\u696d\u7e3e"
            in line
            or "\u696d\u7e3e\uff08"
            in line
            or "\u696d\u7e3e("
            in line
            or "\uff5e"
            in line
        ):
            continue

        if not re.search(
            r"\d{4}\u5e74.*\u671f",
            line,
        ):
            continue

        numbers = re.findall(
            r"-?\d[\d,]*"
            r"(?:\.\d+)?",
            line,
        )

        if len(numbers) < 8:
            continue

        if "\u914d\u5f53" in line:
            continue

        tail = numbers[-8:]

        try:
            yoy_values = [
                float(
                    tail[1].replace(
                        ",",
                        "",
                    )
                ),
                float(
                    tail[3].replace(
                        ",",
                        "",
                    )
                ),
                float(
                    tail[5].replace(
                        ",",
                        "",
                    )
                ),
                float(
                    tail[7].replace(
                        ",",
                        "",
                    )
                ),
            ]
        except Exception:
            continue

        if any(
            abs(value) >= 1000
            for value in yoy_values
        ):
            continue

        candidates.append(
            (
                line,
                numbers,
            )
        )

    if not candidates:
        return None

    return candidates[0]


def extract_yoy(text):
    found = find_performance_row(
        text
    )

    if found is None:
        return None

    line, nums = found
    nums = nums[-8:]

    return {
        "sales_yoy":
            to_float(nums[1]),
        "operating_yoy":
            to_float(nums[3]),
        "ordinary_yoy":
            to_float(nums[5]),
        "net_yoy":
            to_float(nums[7]),
        "source_row":
            line,
    }


def extract_forecast_revision(text):
    patterns = [
        (
            "\u696d\u7e3e\u4e88\u60f3"
            "\u304b\u3089\u306e"
            "\u4fee\u6b63\u306e\u6709\u7121"
        ),
        (
            "\u696d\u7e3e\u4e88\u60f3"
            "\u306e\u4fee\u6b63\u306e\u6709\u7121"
        ),
    ]

    for key in patterns:
        pos = text.find(
            key
        )

        if pos < 0:
            continue

        area = text[
            pos:
            pos + 200
        ]

        if re.search(
            r"[:\uff1a]\s*\u6709",
            area,
        ):
            return True

        if re.search(
            r"[:\uff1a]\s*\u7121",
            area,
        ):
            return False

    return None


def evaluate(result):
    if result is None:
        return "PARSE_ERROR"

    values = [
        result["operating_yoy"],
        result["ordinary_yoy"],
        result["net_yoy"],
    ]

    values = [
        value
        for value in values
        if value is not None
    ]

    if len(values) < 2:
        return "UNKNOWN"

    severe_negative = sum(
        value <= -30.0
        for value in values
    )

    positive = sum(
        value >= 10.0
        for value in values
    )

    strong_positive = sum(
        value >= 30.0
        for value in values
    )

    if severe_negative >= 2:
        return "BAD"

    if (
        strong_positive >= 2
        and min(values) >= 0
    ):
        return "STRONG_GOOD"

    if (
        positive >= 2
        and min(values) > -10
    ):
        return "GOOD"

    return "MIXED"


def format_excel(path):
    wb = load_workbook(
        path
    )

    ws = wb.active
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = (
        ws.dimensions
    )

    for cell in ws[1]:
        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    widths = {
        1: 10,
        2: 24,
        3: 20,
        4: 14,
        5: 14,
        6: 14,
        7: 14,
        8: 16,
        9: 14,
        10: 55,
        11: 80,
        12: 80,
    }

    for col, width in widths.items():
        ws.column_dimensions[
            get_column_letter(col)
        ].width = width

    wb.save(
        path
    )


def main():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--date",
        required=True,
        help="YYYY-MM-DD",
    )

    args = parser.parse_args()

    target_date = pd.Timestamp(
        args.date
    )

    date_api = (
        target_date.strftime(
            "%Y%m%d"
        )
    )

    date_out = (
        target_date.strftime(
            "%Y-%m-%d"
        )
    )

    stocks = load_stock_list()

    universe = set(
        stocks["\u30b3\u30fc\u30c9"]
        .astype(str)
        .str.strip()
    )

    url = (
        "https://webapi.yanoshin.jp/"
        "webapi/tdnet/list/"
        f"{date_api}.json?limit=1000"
    )

    r = requests.get(
        url,
        timeout=30,
        headers={
            "User-Agent":
                "JapanStockScreener/1.0"
        },
    )

    r.raise_for_status()

    data = r.json()

    items = data.get(
        "items",
        [],
    )

    api_count = data.get(
        "total_count",
        None,
    )

    print(
        "=" * 72
    )
    print(
        "EARNINGS CATALYST BATCH"
    )
    print(
        "=" * 72
    )

    print(
        "Date            :",
        date_out,
    )

    print(
        "Stock universe  :",
        len(universe),
    )

    print(
        "API items       :",
        len(items),
    )

    print(
        "API total_count :",
        api_count,
    )

    if len(items) >= 1000:
        print(
            "WARNING: API may be truncated"
        )

    rows = []

    targets = 0
    parse_errors = 0
    download_errors = 0

    cutoff = (
        target_date
        + pd.Timedelta(
            hours=15,
            minutes=30,
        )
    )

    for item in items:
        tdnet = item.get(
            "Tdnet",
            item,
        )

        code = normalize_code(
            tdnet.get(
                "company_code",
                "",
            )
        )

        if code not in universe:
            continue

        title = str(
            tdnet.get(
                "title",
                "",
            )
        ).strip()

        if (
            "\u6c7a\u7b97\u77ed\u4fe1"
            not in title
        ):
            continue

        if "\u8a02\u6b63" in title:
            continue

        title_upper = title.upper()

        if (
            "IFRS"
            in title_upper
            or "\uff29\uff26\uff32\uff33"
            in title
        ):
            continue

        pubdate = str(
            tdnet.get(
                "pubdate",
                "",
            )
        ).strip()

        try:
            pub_dt = pd.to_datetime(
                pubdate
            )
        except Exception:
            continue

        if pub_dt < cutoff:
            continue

        targets += 1

        name = str(
            tdnet.get(
                "company_name",
                "",
            )
        ).strip()

        pdf_url = str(
            tdnet.get(
                "document_url",
                "",
            )
        ).strip()

        row = {
            "\u30b3\u30fc\u30c9":
                code,
            "\u9298\u67c4\u540d":
                name,
            "\u767a\u8868\u65e5\u6642":
                pubdate,
            "\u58f2\u4e0a\u524d\u5e74\u6bd4":
                None,
            "\u55b6\u696d\u5229\u76ca\u524d\u5e74\u6bd4":
                None,
            "\u7d4c\u5e38\u5229\u76ca\u524d\u5e74\u6bd4":
                None,
            "\u7d14\u5229\u76ca\u524d\u5e74\u6bd4":
                None,
            "\u6c7a\u7b97\u8a55\u4fa1":
                "",
            "\u696d\u7e3e\u4e88\u60f3\u4fee\u6b63":
                "",
            "\u958b\u793a\u30bf\u30a4\u30c8\u30eb":
                title,
            "\u62bd\u51fa\u5143\u884c":
                "",
            "\u958b\u793aURL":
                pdf_url,
        }

        try:
            pdf_path = download_pdf(
                pdf_url,
                date_api,
                code,
            )

            text = extract_first_page(
                pdf_path
            )

            result = extract_yoy(
                text
            )

            revision = (
                extract_forecast_revision(
                    text
                )
            )

            if result is None:
                row[
                    "\u6c7a\u7b97\u8a55\u4fa1"
                ] = "PARSE_ERROR"

                parse_errors += 1

            else:
                row[
                    "\u58f2\u4e0a\u524d\u5e74\u6bd4"
                ] = result[
                    "sales_yoy"
                ]

                row[
                    "\u55b6\u696d\u5229\u76ca\u524d\u5e74\u6bd4"
                ] = result[
                    "operating_yoy"
                ]

                row[
                    "\u7d4c\u5e38\u5229\u76ca\u524d\u5e74\u6bd4"
                ] = result[
                    "ordinary_yoy"
                ]

                row[
                    "\u7d14\u5229\u76ca\u524d\u5e74\u6bd4"
                ] = result[
                    "net_yoy"
                ]

                row[
                    "\u6c7a\u7b97\u8a55\u4fa1"
                ] = evaluate(
                    result
                )

                row[
                    "\u62bd\u51fa\u5143\u884c"
                ] = result[
                    "source_row"
                ]

            if revision is True:
                row[
                    "\u696d\u7e3e\u4e88\u60f3\u4fee\u6b63"
                ] = "\u6709"

            elif revision is False:
                row[
                    "\u696d\u7e3e\u4e88\u60f3\u4fee\u6b63"
                ] = "\u7121"

            else:
                row[
                    "\u696d\u7e3e\u4e88\u60f3\u4fee\u6b63"
                ] = "\u4e0d\u660e"

        except Exception as e:
            row[
                "\u6c7a\u7b97\u8a55\u4fa1"
            ] = "DOWNLOAD_ERROR"

            row[
                "\u62bd\u51fa\u5143\u884c"
            ] = repr(e)

            download_errors += 1

        rows.append(
            row
        )

    df = pd.DataFrame(
        rows
    )

    if not df.empty:
        priority = {
            "STRONG_GOOD": 0,
            "GOOD": 1,
            "MIXED": 2,
            "BAD": 3,
            "UNKNOWN": 4,
            "PARSE_ERROR": 5,
            "DOWNLOAD_ERROR": 6,
        }

        df["_sort"] = (
            df[
                "\u6c7a\u7b97\u8a55\u4fa1"
            ]
            .map(priority)
            .fillna(99)
        )

        df = (
            df.sort_values(
                [
                    "_sort",
                    "\u55b6\u696d\u5229\u76ca\u524d\u5e74\u6bd4",
                ],
                ascending=[
                    True,
                    False,
                ],
                na_position="last",
            )
            .drop(
                columns=[
                    "_sort"
                ]
            )
        )

    out = (
        Path("results")
        / (
            f"{date_out}_"
            "earnings_catalyst.xlsx"
        )
    )

    out.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    df.to_excel(
        out,
        index=False,
        sheet_name=
            "\u6c7a\u7b97\u89e3\u6790",
    )

    format_excel(
        out
    )

    print()
    print(
        "Eligible earnings :",
        targets,
    )

    print(
        "Parsed            :",
        targets
        - parse_errors
        - download_errors,
    )

    print(
        "Parse errors      :",
        parse_errors,
    )

    print(
        "Download errors   :",
        download_errors,
    )

    if not df.empty:
        print()
        print(
            df[
                "\u6c7a\u7b97\u8a55\u4fa1"
            ]
            .value_counts(
                dropna=False
            )
            .to_string()
        )

    print()
    print(
        "Saved :",
        out.resolve(),
    )


if __name__ == "__main__":
    main()
