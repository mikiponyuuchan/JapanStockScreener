from __future__ import annotations

import argparse
from datetime import datetime
from pathlib import Path

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import sys

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"

if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

from screener.loader import load_stock_list


RESULTS_DIR = ROOT / "results"
API_URL = (
    "https://webapi.yanoshin.jp/"
    "webapi/tdnet/list/{date}.json"
)

TIMEOUT = 30


RULES = [
    (
        "TOB/MBO",
        5,
        [
            "\u516c\u958b\u8cb7\u4ed8\u3051\u306e\u958b\u59cb",
            "\u516c\u958b\u8cb7\u4ed8\u306e\u958b\u59cb",
            "\u516c\u958b\u8cb7\u4ed8\u3051\u306b\u95a2\u3059\u308b\u8cdb\u540c",
            "\u516c\u958b\u8cb7\u4ed8\u306b\u95a2\u3059\u308b\u8cdb\u540c",
            "MBO",
            "\u30de\u30cd\u30b8\u30e1\u30f3\u30c8\u30fb\u30d0\u30a4\u30a2\u30a6\u30c8",
        ],
        False,
    ),
    (
        "\u4e0a\u65b9\u4fee\u6b63",
        4,
        [
            "\u4e0a\u65b9\u4fee\u6b63",
            "\u696d\u7e3e\u4e88\u60f3\u306e\u4e0a\u65b9\u4fee\u6b63",
        ],
        False,
    ),
    (
        "\u9ed2\u5b57\u8ee2\u63db",
        4,
        [
            "\u9ed2\u5b57\u8ee2\u63db",
            "\u9ed2\u5b57\u5316",
        ],
        False,
    ),
    (
        "\u81ea\u5df1\u682a\u5f0f\u53d6\u5f97",
        4,
        [
            "\u81ea\u5df1\u682a\u5f0f\u306e\u53d6\u5f97\u306b\u4fc2\u308b\u4e8b\u9805\u306e\u6c7a\u5b9a",
            "\u81ea\u5df1\u682a\u5f0f\u53d6\u5f97\u306b\u4fc2\u308b\u4e8b\u9805\u306e\u6c7a\u5b9a",
            "\u81ea\u5df1\u682a\u5f0f\u306e\u53d6\u5f97\u53ca\u3073\u6d88\u5374\u306b\u4fc2\u308b\u4e8b\u9805\u306e\u6c7a\u5b9a",
        ],
        False,
    ),
    (
        "\u682a\u4e3b\u512a\u5f85",
        4,
        [
            "\u682a\u4e3b\u512a\u5f85\u5236\u5ea6\u306e\u65b0\u8a2d",
            "\u682a\u4e3b\u512a\u5f85\u306e\u65b0\u8a2d",
            "\u682a\u4e3b\u512a\u5f85\u5236\u5ea6\u306e\u62e1\u5145",
            "\u682a\u4e3b\u512a\u5f85\u306e\u62e1\u5145",
        ],
        False,
    ),
    (
        "\u5897\u914d",
        3,
        [
            "\u5897\u914d",
            "\u7279\u5225\u914d\u5f53",
            "\u8a18\u5ff5\u914d\u5f53",
        ],
        False,
    ),
    (
        "\u5927\u578b\u53d7\u6ce8",
        3,
        [
            "\u5927\u53e3\u53d7\u6ce8",
            "\u5927\u578b\u53d7\u6ce8",
            "\u53d7\u6ce8\u306b\u95a2\u3059\u308b\u304a\u77e5\u3089\u305b",
        ],
        False,
    ),
    (
        "\u8cc7\u672c\u696d\u52d9\u63d0\u643a",
        3,
        [
            "\u8cc7\u672c\u696d\u52d9\u63d0\u643a",
            "\u8cc7\u672c\u63d0\u643a",
        ],
        False,
    ),
    (
        "\u627f\u8a8d",
        3,
        [
            "\u627f\u8a8d\u53d6\u5f97",
            "\u88fd\u9020\u8ca9\u58f2\u627f\u8a8d",
        ],
        False,
    ),
    (
        "\u4e0b\u65b9\u4fee\u6b63",
        0,
        [
            "\u4e0b\u65b9\u4fee\u6b63",
            "\u696d\u7e3e\u4e88\u60f3\u306e\u4e0b\u65b9\u4fee\u6b63",
        ],
        True,
    ),
    (
        "\u8d64\u5b57\u8ee2\u843d",
        0,
        [
            "\u8d64\u5b57\u8ee2\u843d",
            "\u8d64\u5b57\u898b\u8fbc",
        ],
        True,
    ),
    (
        "\u5e0c\u8584\u5316",
        0,
        [
            "\u7b2c\u4e09\u8005\u5272\u5f53\u306b\u3088\u308b\u65b0\u682a\u5f0f",
            "\u7b2c\u4e09\u8005\u5272\u5f53\u5897\u8cc7",
            "\u65b0\u682a\u4e88\u7d04\u6a29\u306e\u767a\u884c",
        ],
        True,
    ),
]


def normalize_code(value) -> str:
    text = str(value).strip()

    if text.endswith(".0"):
        text = text[:-2]

    if len(text) == 5 and text.endswith("0"):
        text = text[:-1]

    return text


def find_column(df: pd.DataFrame, candidates: list[str]):
    for candidate in candidates:
        if candidate in df.columns:
            return candidate
    return None


def build_stock_map():
    stocks = load_stock_list()

    code_col = find_column(
        stocks,
        [
            "\u30b3\u30fc\u30c9",
            "Code",
            "code",
        ],
    )

    name_col = find_column(
        stocks,
        [
            "\u9285\u67c4\u540d",
            "\u4f1a\u793e\u540d",
            "Name",
            "name",
        ],
    )

    if code_col is None:
        raise RuntimeError(
            "Stock code column was not found."
        )

    stock_map = {}

    for _, row in stocks.iterrows():
        code = normalize_code(
            row[code_col]
        )

        if not code:
            continue

        name = ""

        if name_col is not None:
            value = row[name_col]
            if pd.notna(value):
                name = str(value).strip()

        stock_map[code] = name

    return stock_map


def download_tdnet(date_text: str):
    api_date = date_text.replace("-", "")
    url = API_URL.format(date=api_date)

    response = requests.get(
        url,
        timeout=TIMEOUT,
        headers={
            "User-Agent":
                "JapanStockScreener/1.0"
        },
    )

    response.raise_for_status()

    try:
        data = response.json()
    except Exception:
        import json
        data = json.loads(response.text)

    return data, url


def classify_title(title: str):
    matched = []

    for category, importance, words, avoid in RULES:
        if any(
            word.lower() in title.lower()
            for word in words
        ):
            matched.append(
                {
                    "category": category,
                    "importance": importance,
                    "avoid": avoid,
                }
            )

    if not matched:
        return None

    avoid_matches = [
        item for item in matched
        if item["avoid"]
    ]

    positive_matches = [
        item for item in matched
        if not item["avoid"]
    ]

    if avoid_matches:
        chosen = avoid_matches[0]
    else:
        chosen = max(
            positive_matches,
            key=lambda x: x["importance"],
        )

    categories = " / ".join(
        dict.fromkeys(
            item["category"]
            for item in matched
        )
    )

    return {
        "category": categories,
        "importance": chosen["importance"],
        "avoid": bool(avoid_matches),
    }


def importance_text(value: int, avoid: bool):
    if avoid:
        return "\u56de\u907f"

    return "\u2605" * int(value)


def attention_text(value: int, avoid: bool):
    if avoid:
        return "\u56de\u907f"

    if value >= 5:
        return "\u6700\u512a\u5148"

    if value >= 4:
        return "\u9ad8"

    return "\u6ce8\u76ee"


def create_dataframe(
    items,
    stock_map,
):
    rows = []

    for item in items:
        tdnet = item.get(
            "Tdnet",
            item,
        )

        raw_code = tdnet.get(
            "company_code",
            "",
        )

        code = normalize_code(raw_code)

        if code not in stock_map:
            continue

        pubdate = str(
            tdnet.get(
                "pubdate",
                "",
            )
        ).strip()

        try:
            pub_dt = pd.to_datetime(
                pubdate
            )
        except Exception:
            continue

        cutoff_time = pd.Timestamp(
            pub_dt.date()
        ) + pd.Timedelta(
            hours=15,
            minutes=30,
        )

        if pub_dt < cutoff_time:
            continue

        title = str(
            tdnet.get(
                "title",
                "",
            )
        ).strip()

        tob_noise_words = [
            "\u516c\u958b\u8cb7\u4ed8\u3051\u306e\u7d50\u679c",
            "\u516c\u958b\u8cb7\u4ed8\u306e\u7d50\u679c",
            "\u516c\u958b\u8cb7\u4ed8\u3051\u306e\u5909\u66f4",
            "\u516c\u958b\u8cb7\u4ed8\u306e\u5909\u66f4",
            "\u516c\u958b\u8cb7\u4ed8\u5c4a\u51fa\u66f8\u306e\u8a02\u6b63",
        ]

        if any(
            word in title
            for word in tob_noise_words
        ):
            continue

        classification = classify_title(
            title
        )

        if classification is None:
            continue

        company_name = str(
            tdnet.get(
                "company_name",
                "",
            )
        ).strip()

        if not company_name:
            company_name = stock_map.get(
                code,
                "",
            )

        importance = classification[
            "importance"
        ]

        avoid = classification[
            "avoid"
        ]

        rows.append(
            {
                "\u30b3\u30fc\u30c9":
                    code,
                "\u9285\u67c4\u540d":
                    company_name,
                "\u767a\u8868\u65e5\u6642":
                    tdnet.get(
                        "pubdate",
                        "",
                    ),
                "\u6750\u6599\u5206\u985e":
                    classification[
                        "category"
                    ],
                "\u91cd\u8981\u5ea6":
                    importance_text(
                        importance,
                        avoid,
                    ),
                "\u7fcc\u671d\u6ce8\u76ee":
                    attention_text(
                        importance,
                        avoid,
                    ),
                "\u56de\u907f\u30d5\u30e9\u30b0":
                    "\u25cb"
                    if avoid
                    else "",
                "\u958b\u793a\u30bf\u30a4\u30c8\u30eb":
                    title,
                "\u958b\u793aURL":
                    tdnet.get(
                        "document_url",
                        "",
                    ),
                "_importance":
                    importance,
                "_avoid":
                    int(avoid),
            }
        )

    if not rows:
        return pd.DataFrame(
            columns=[
                "\u30b3\u30fc\u30c9",
                "\u9285\u67c4\u540d",
                "\u767a\u8868\u65e5\u6642",
                "\u6750\u6599\u5206\u985e",
                "\u91cd\u8981\u5ea6",
                "\u7fcc\u671d\u6ce8\u76ee",
                "\u56de\u907f\u30d5\u30e9\u30b0",
                "\u958b\u793a\u30bf\u30a4\u30c8\u30eb",
                "\u958b\u793aURL",
            ]
        )

    df = pd.DataFrame(rows)

    df = df.sort_values(
        by=[
            "_avoid",
            "_importance",
            "\u767a\u8868\u65e5\u6642",
        ],
        ascending=[
            True,
            False,
            False,
        ],
    )

    df = df.drop(
        columns=[
            "_importance",
            "_avoid",
        ]
    )

    return df.reset_index(
        drop=True
    )


def format_excel(path: Path):
    wb = load_workbook(path)
    ws = wb.active
    ws.title = "\u7fcc\u671d\u6750\u6599\u5019\u88dc"

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="D9EAF7",
    )

    for cell in ws[1]:
        cell.font = Font(
            bold=True
        )
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    widths = {
        1: 11,
        2: 24,
        3: 20,
        4: 24,
        5: 12,
        6: 12,
        7: 12,
        8: 70,
        9: 55,
    }

    for col_idx, width in widths.items():
        ws.column_dimensions[
            get_column_letter(col_idx)
        ].width = width

    for row in ws.iter_rows(
        min_row=2
    ):
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    url_col = 9

    for row_idx in range(
        2,
        ws.max_row + 1,
    ):
        cell = ws.cell(
            row=row_idx,
            column=url_col,
        )

        if cell.value:
            cell.hyperlink = str(
                cell.value
            )
            cell.style = "Hyperlink"

    wb.save(path)


def parse_args():
    parser = argparse.ArgumentParser()

    parser.add_argument(
        "--date",
        default=datetime.now().strftime(
            "%Y-%m-%d"
        ),
        help="TDnet date YYYY-MM-DD",
    )

    return parser.parse_args()


def main():
    args = parse_args()
    date_text = args.date

    print("=" * 72)
    print("MORNING CATALYST Ver1")
    print("=" * 72)
    print(
        f"TDnet date : {date_text}"
    )

    stock_map = build_stock_map()

    print(
        f"Stock universe : "
        f"{len(stock_map)}"
    )

    data, api_url = download_tdnet(
        date_text
    )

    items = data.get(
        "items",
        [],
    )

    print(
        f"TDnet disclosures : "
        f"{len(items)}"
    )

    df = create_dataframe(
        items,
        stock_map,
    )

    RESULTS_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    output_path = (
        RESULTS_DIR
        / (
            f"{date_text}"
            "_morning_catalyst.xlsx"
        )
    )

    df.to_excel(
        output_path,
        index=False,
        engine="openpyxl",
    )

    format_excel(
        output_path
    )

    print(
        f"Catalyst candidates : "
        f"{len(df)}"
    )

    if not df.empty:
        display_cols = [
            "\u30b3\u30fc\u30c9",
            "\u9285\u67c4\u540d",
            "\u767a\u8868\u65e5\u6642",
            "\u6750\u6599\u5206\u985e",
            "\u91cd\u8981\u5ea6",
            "\u7fcc\u671d\u6ce8\u76ee",
            "\u958b\u793a\u30bf\u30a4\u30c8\u30eb",
        ]

        print()
        print(
            df[
                display_cols
            ].to_string(
                index=False
            )
        )

    print()
    print(
        f"Saved : {output_path}"
    )
    print(
        f"Source: {api_url}"
    )


if __name__ == "__main__":
    main()
